"""Validation of generated PPTX files.

A file that ends in ``.pptx`` is not automatically an editable presentation: the
easiest failure mode in this problem space is a deck made of slide-sized images.
These checks look inside the OOXML package and confirm that text, charts and
tables are native objects, that the ZIP is intact, and that Arabic decks carry
paragraph-level RTL markers.

This is the same validator used by the benchmark tool and by the request path,
so a deck served to a student has passed exactly the checks the benchmark
reports.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

_SLIDE_RE = re.compile(r"^ppt/slides/slide\d+\.xml$")
_CHART_RE = re.compile(r"^ppt/charts/chart\d+\.xml$")


class DeckValidationError(ValueError):
    """The generated file is not a valid, editable presentation."""


@dataclass(frozen=True, slots=True)
class DeckValidation:
    size_bytes: int
    slides: int
    text_shapes: int
    chart_parts: int
    table_parts: int
    rtl_paragraphs: int
    table_rtl_parts: int
    zip_ok: bool
    editable: bool
    issues: tuple[str, ...]

    @property
    def ok(self) -> bool:
        return self.zip_ok and self.editable and not self.issues


def inspect(path: Path) -> DeckValidation:
    """Inspect a PPTX package without trusting its file extension."""
    issues: list[str] = []
    size = path.stat().st_size
    try:
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            if bad is not None:
                issues.append(f"corrupt zip member: {bad}")
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
                issues.append("missing core OOXML parts")
            slides = sorted(name for name in names if _SLIDE_RE.match(name))
            chart_parts = sorted(name for name in names if _CHART_RE.match(name))
            table_parts = 0
            rtl_paragraphs = 0
            table_rtl_parts = 0
            text_shapes = 0
            for slide_name in slides:
                xml = archive.read(slide_name).decode("utf-8", errors="replace")
                text_shapes += xml.count("<p:sp>")
                table_parts += xml.count("<a:tbl>")
                rtl_paragraphs += len(re.findall(r'<a:pPr[^>]*rtl="1"', xml))
                table_rtl_parts += len(re.findall(r'<a:tblPr[^>]*rtl="1"', xml))
            if not slides:
                issues.append("no slides found")
    except (zipfile.BadZipFile, OSError) as exc:
        raise DeckValidationError("generated file is not a readable PPTX package") from exc

    # Opening with python-pptx proves the package is well-formed enough for a
    # real consumer, and lets us count *native* shapes rather than XML tags.
    editable = False
    try:
        from pptx import Presentation  # local import keeps the module importable

        presentation = Presentation(str(path))
        native = 0
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    native += 1
                if getattr(shape, "has_chart", False):
                    native += 1
                if getattr(shape, "has_table", False):
                    native += 1
        editable = native > 0 and len(presentation.slides) == len(slides)
        if not editable:
            issues.append("no editable native shapes found")
    except Exception as exc:
        raise DeckValidationError("generated file cannot be opened as a presentation") from exc

    return DeckValidation(
        size_bytes=size,
        slides=len(slides),
        text_shapes=text_shapes,
        chart_parts=len(chart_parts),
        table_parts=table_parts,
        rtl_paragraphs=rtl_paragraphs,
        table_rtl_parts=table_rtl_parts,
        zip_ok="corrupt zip member" not in " ".join(issues),
        editable=editable,
        issues=tuple(issues),
    )


def validate(
    path: Path,
    *,
    expected_slides: int,
    require_rtl: bool,
    require_chart: bool = False,
    require_table: bool = False,
    max_bytes: int | None = None,
) -> DeckValidation:
    """Raise ``DeckValidationError`` when a gate fails; otherwise return evidence."""
    report = inspect(path)
    problems: list[str] = list(report.issues)
    if report.slides != expected_slides:
        problems.append(f"expected {expected_slides} slides, found {report.slides}")
    if require_rtl and report.rtl_paragraphs == 0:
        problems.append("no RTL paragraph markers in an Arabic deck")
    if require_chart and report.chart_parts == 0:
        problems.append("chart slide produced no chart part")
    if require_table and report.table_parts == 0:
        problems.append("table slide produced no table")
    if max_bytes is not None and report.size_bytes > max_bytes:
        problems.append("generated deck exceeds the output size cap")
    if not report.editable or not report.zip_ok:
        problems.append("generated deck is not an intact, editable PPTX")
    if problems:
        raise DeckValidationError("; ".join(problems))
    return report


__all__ = ["PPTX_MEDIA_TYPE", "DeckValidation", "DeckValidationError", "inspect", "validate"]
