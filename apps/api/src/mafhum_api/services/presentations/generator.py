"""Editable PPTX generation.

The production generator uses python-pptx (MIT). It writes **native PowerPoint
objects** — text boxes, a real chart and a real table — never a rendered picture
of a slide. Arabic direction is written at the paragraph level, which is the
unit PowerPoint actually uses: ``a:pPr@rtl="1"`` plus right alignment and an
Arabic complex-script typeface and language on every run.

Text is stored in logical order. Reversing Arabic glyphs or injecting bidi
control characters here would corrupt mixed Arabic/French citations once the
student edits them, so we let the renderer apply the Unicode bidi algorithm.

A ``DeckGenerator`` protocol keeps the service independent of the library: the
Node/PptxGenJS path used by the benchmark implements the same spec contract.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol, runtime_checkable

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from mafhum_api.services.presentations.spec import DeckSlide, DeckSpec

NAVY = RGBColor(0x0B, 0x25, 0x45)  # type: ignore[no-untyped-call]
ACCENT = RGBColor(0x1B, 0x6C, 0xA8)  # type: ignore[no-untyped-call]
PALE = RGBColor(0xF1, 0xF5, 0xFA)  # type: ignore[no-untyped-call]
WHITE = RGBColor(0xFF, 0xFF, 0xFF)  # type: ignore[no-untyped-call]
INK = RGBColor(0x14, 0x1B, 0x2D)  # type: ignore[no-untyped-call]
MUTED = RGBColor(0x4A, 0x5A, 0x70)  # type: ignore[no-untyped-call]

# A widely available font. Nothing is embedded: PowerPoint/LibreOffice performs
# substitution where the font is missing, and the deck remains fully editable.
ARABIC_FONT = "Arial"
LATIN_FONT = "Arial"
ARABIC_LANG = "ar-DZ"
FRENCH_LANG = "fr-FR"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


@runtime_checkable
class DeckGenerator(Protocol):
    """Generate one PPTX file from a validated spec."""

    name: str
    version: str

    def generate(self, spec: DeckSpec, output: Path) -> None: ...


@dataclass(frozen=True, slots=True)
class GeneratedDeck:
    path: Path
    size_bytes: int
    sha256: str


def describe_file(path: Path) -> GeneratedDeck:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 256), b""):
            digest.update(chunk)
            size += len(chunk)
    return GeneratedDeck(path=path, size_bytes=size, sha256=digest.hexdigest())


def _set_run_metadata(run: object, *, rtl: bool) -> None:
    """Set language and a complex-script typeface on one run's ``a:rPr``.

    ``a:latin`` alone is not enough for Arabic: PowerPoint resolves Arabic
    characters through the complex-script font (``a:cs``). Without it, Arabic
    text can render with a fallback that breaks shaping and editability.
    """
    r_pr = run._r.get_or_add_rPr()  # type: ignore[attr-defined]
    r_pr.set("lang", ARABIC_LANG if rtl else FRENCH_LANG)
    typeface = ARABIC_FONT if rtl else LATIN_FONT
    for tag in ("a:cs", "a:ea"):
        element = r_pr.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            r_pr.append(element)
        element.set("typeface", typeface)


def _style_paragraph(paragraph: object, *, rtl: bool) -> None:
    p_pr = paragraph._p.get_or_add_pPr()  # type: ignore[attr-defined]
    p_pr.set("rtl", "1" if rtl else "0")
    p_pr.set("algn", "r" if rtl else "l")
    paragraph.alignment = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT  # type: ignore[attr-defined]
    for run in paragraph.runs:  # type: ignore[attr-defined]
        _set_run_metadata(run, rtl=rtl)


def _textbox(
    slide: object,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    text: str,
    rtl: bool,
    size: int,
    color: RGBColor,
    bold: bool = False,
    fill: RGBColor | None = None,
    margin: float = 0.12,
) -> Any:
    shape = slide.shapes.add_textbox(left, top, width, height)  # type: ignore[attr-defined]
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(0.05)
    frame.margin_bottom = Inches(0.05)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    lines = text.split("\n") or [""]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.name = LATIN_FONT
        paragraph.font.color.rgb = color
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = LATIN_FONT
            run.font.color.rgb = color
            _set_run_metadata(run, rtl=rtl)
        _style_paragraph(paragraph, rtl=rtl)
    return shape


def _add_band(slide: object, *, rtl: bool, title: str, draft_notice: str, slide_no: int) -> None:
    band = slide.shapes.add_shape(  # type: ignore[attr-defined]
        1,  # MSO_SHAPE.RECTANGLE
        0,
        0,
        SLIDE_W,
        Inches(1.15),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    _textbox(
        slide,
        left=Inches(0.55),
        top=Inches(0.12),
        width=SLIDE_W - Inches(1.1),
        height=Inches(0.9),
        text=title,
        rtl=rtl,
        size=26,
        color=WHITE,
        bold=True,
    )
    footer = f"{draft_notice}  •  MAFHUM  •  {slide_no}"
    _textbox(
        slide,
        left=Inches(0.55),
        top=SLIDE_H - Inches(0.48),
        width=SLIDE_W - Inches(1.1),
        height=Inches(0.32),
        text=footer,
        rtl=rtl,
        size=10,
        color=MUTED,
    )


def _citation_footer(slide: object, *, rtl: bool, indices: tuple[int, ...]) -> None:
    if not indices:
        return
    label = "[" + "، ".join(str(i) for i in indices) + "]"
    _textbox(
        slide,
        left=Inches(0.55),
        top=SLIDE_H - Inches(0.95),
        width=SLIDE_W - Inches(1.1),
        height=Inches(0.35),
        text=label,
        rtl=rtl,
        size=11,
        color=ACCENT,
        bold=True,
    )


class PythonPptxGenerator:
    """Write an editable, RTL-correct deck using python-pptx."""

    name = "python-pptx"
    version = "1.0.2"

    def generate(self, spec: DeckSpec, output: Path) -> None:
        spec.validate()
        presentation = Presentation()
        presentation.slide_width = SLIDE_W
        presentation.slide_height = SLIDE_H
        blank = presentation.slide_layouts[6]

        for number, deck_slide in enumerate(spec.slides, start=1):
            slide = presentation.slides.add_slide(blank)
            if deck_slide.kind != "title":
                _add_band(
                    slide,
                    rtl=spec.rtl,
                    title=deck_slide.title,
                    draft_notice=spec.draft_notice,
                    slide_no=number,
                )
            self._render_slide(slide, deck_slide, spec)
            _citation_footer(slide, rtl=spec.rtl, indices=deck_slide.citation_indices)

        output.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(output))

    def _render_slide(self, slide: object, deck_slide: DeckSlide, spec: DeckSpec) -> None:
        rtl = spec.rtl
        if deck_slide.kind == "title":
            background = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)  # type: ignore[attr-defined]
            background.fill.solid()
            background.fill.fore_color.rgb = NAVY
            background.line.fill.background()
            _textbox(
                slide,
                left=Inches(0.8),
                top=Inches(2.1),
                width=SLIDE_W - Inches(1.6),
                height=Inches(1.4),
                text=spec.title,
                rtl=rtl,
                size=40,
                color=WHITE,
                bold=True,
            )
            _textbox(
                slide,
                left=Inches(0.8),
                top=Inches(3.7),
                width=SLIDE_W - Inches(1.6),
                height=Inches(0.8),
                text=spec.subtitle,
                rtl=rtl,
                size=20,
                color=PALE,
            )
            _textbox(
                slide,
                left=Inches(0.8),
                top=Inches(5.6),
                width=SLIDE_W - Inches(1.6),
                height=Inches(0.6),
                text=spec.draft_notice,
                rtl=rtl,
                size=14,
                color=PALE,
            )
            return

        top = Inches(1.45)
        width = SLIDE_W - Inches(1.1)
        if deck_slide.kind == "bullets":
            _textbox(
                slide,
                left=Inches(0.55),
                top=top,
                width=width,
                height=Inches(4.6),
                text="\n".join(f"• {item}" for item in deck_slide.items),
                rtl=rtl,
                size=18,
                color=INK,
            )
            return
        if deck_slide.kind == "quote":
            _textbox(
                slide,
                left=Inches(0.55),
                top=top,
                width=width,
                height=Inches(3.1),
                text=deck_slide.text,
                rtl=rtl,
                size=18,
                color=INK,
                fill=PALE,
            )
            _textbox(
                slide,
                left=Inches(0.55),
                top=Inches(4.85),
                width=width,
                height=Inches(0.7),
                text=deck_slide.items[0] if deck_slide.items else "",
                rtl=rtl,
                size=13,
                color=MUTED,
            )
            return
        if deck_slide.kind == "chart":
            data = CategoryChartData()  # type: ignore[no-untyped-call]
            data.categories = list(deck_slide.chart_labels)
            data.add_series(  # type: ignore[no-untyped-call]
                "citations" if not rtl else "الاقتباسات", tuple(deck_slide.chart_values)
            )
            chart = slide.shapes.add_chart(  # type: ignore[attr-defined]
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                Inches(0.7),
                top,
                width - Inches(0.3),
                Inches(4.5),
                data,
            ).chart
            chart.has_legend = False
            return
        if deck_slide.kind == "table":
            rows = max(2, len(deck_slide.table_rows) + 1)
            cols = max(1, len(deck_slide.table_headers))
            table_shape = slide.shapes.add_table(  # type: ignore[attr-defined]
                rows, cols, Inches(0.55), top, width, Inches(4.4)
            )
            table = table_shape.table
            # Table-level direction is separate from paragraph direction. Without
            # a:tblPr@rtl, PowerPoint may lay columns out left-to-right even
            # though every cell's text is right-aligned.
            tbl_pr = table._tbl.find(qn("a:tblPr"))
            if tbl_pr is not None:
                tbl_pr.set("rtl", "1" if rtl else "0")
            for col_index, header in enumerate(deck_slide.table_headers):
                cell = table.cell(0, col_index)
                cell.text = header
                for paragraph in cell.text_frame.paragraphs:
                    _style_paragraph(paragraph, rtl=rtl)
            for row_index, values in enumerate(deck_slide.table_rows, start=1):
                for col_index, value in enumerate(values):
                    cell = table.cell(row_index, col_index)
                    cell.text = value
                    for paragraph in cell.text_frame.paragraphs:
                        _style_paragraph(paragraph, rtl=rtl)


__all__ = [
    "DeckGenerator",
    "GeneratedDeck",
    "PythonPptxGenerator",
    "describe_file",
]
